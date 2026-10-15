"""Build the IC, manager, and director resume editions.

Parses the canonical Markdown resume (carlos-barajas-resume.md), keeps the
Experience bullets tagged for each edition, renders each edition to .docx
with python-docx following docs/style-guide.md, converts the .docx to PDF
with Microsoft Word, and writes the PDFs to public/resumes/ (served by the
site and deployed to the Static Web App). It also writes
src/data/resume-editions.json, which the site reads for its download links
and a unit test uses to catch PDFs that are stale relative to the Markdown.

Run from the repo root with `pnpm resumes`. Requires macOS, Microsoft Word,
and uv. See scripts/resume/README.md.
"""

from __future__ import annotations

import hashlib
import json
import re
import shutil
import subprocess
import sys
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_LINE_SPACING
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor
from pypdf import PdfReader

REPO = Path(__file__).resolve().parents[2]
SOURCE = REPO / "carlos-barajas-resume.md"
PDF_DIR = REPO / "public" / "resumes"
MANIFEST = REPO / "src" / "data" / "resume-editions.json"
DOCX_DIR = REPO / "scripts" / "resume" / "out"
# Word is sandboxed; its own container is the one place it can read and
# write without a file-access prompt, so conversion happens there.
WORD_WORKDIR = Path.home() / "Library/Containers/com.microsoft.Word/Data/tmp/profilesite-resumes"

FILE_SLUGS = {"ic": "ic", "mgr": "manager", "dir": "director"}

# Tokens from docs/style-guide.md. PDFs use Word-bundled fonts: Century
# Gothic stands in for Montserrat, Aptos for Lato.
DISPLAY_FONT = "Century Gothic"
BODY_FONT = "Aptos"
INK = "1A1D21"
MUTED = "4F5B66"
BLUE = "0076B9"
BLOCK = "C3E2F4"
BLOCK_SOFT = "DBEAF1"
BAND = "EDF2F6"
RULE = "D7D8DA"
RULE_STRONG = "CBCCCE"

PAGE_WIDTH = Inches(8.5)
PAGE_HEIGHT = Inches(11)
MARGIN = Inches(0.5)
CONTENT_WIDTH = PAGE_WIDTH - 2 * MARGIN
MAIN_WIDTH = Inches(5.35)
MAIN_GUTTER = Pt(14)
MAIN_CONTENT_WIDTH = MAIN_WIDTH - MAIN_GUTTER


# ---------------------------------------------------------------------------
# Markdown model and parser
# ---------------------------------------------------------------------------


@dataclass
class Bullet:
    text: str
    audiences: list[str]
    sub_items: list[str] = field(default_factory=list)


@dataclass
class Role:
    title: str
    range: str
    bullets: list[Bullet] = field(default_factory=list)


@dataclass
class Company:
    name: str
    range: str
    note: str | None
    roles: list[Role] = field(default_factory=list)


@dataclass
class EarlierRole:
    company: str
    title: str
    range: str
    description: str


@dataclass
class Variant:
    id: str
    label: str
    headline: str
    summary: str


@dataclass
class Resume:
    name: str
    contact: dict[str, str]
    skills: list[tuple[str, list[str]]]
    experience: list[Company]
    earlier: list[EarlierRole]
    education: tuple[str, str]
    interests: list[str]
    variants: list[Variant]


def split_items(items: str) -> list[str]:
    """Split a comma-separated skills line, ignoring commas inside parentheses."""
    parts, depth, current = [], 0, ""
    for char in items:
        depth += char == "("
        depth -= char == ")"
        if char == "," and depth == 0:
            parts.append(current.strip())
            current = ""
        else:
            current += char
    return [*parts, current.strip()]


def sections(markdown: str) -> dict[str, str]:
    chunks = re.split(r"^## (.+)$", markdown, flags=re.M)
    return {"_header": chunks[0], **dict(zip(chunks[1::2], chunks[2::2], strict=True))}


def subsections(body: str, level: str) -> list[tuple[str, str]]:
    chunks = re.split(rf"^{level} (.+)$", body, flags=re.M)
    return list(zip(chunks[1::2], chunks[2::2], strict=True))


def parse_bullets(body: str) -> list[Bullet]:
    bullets: list[Bullet] = []
    for line in body.splitlines():
        if match := re.match(r"^- (.+) <!-- ([a-z ]+) -->$", line):
            bullets.append(Bullet(match[1], match[2].split()))
        elif match := re.match(r"^  - (.+)$", line):
            bullets[-1].sub_items.append(match[1])
    return bullets


def parse(markdown: str) -> Resume:
    # Drop multi-line comment blocks; per-bullet audience tags stay because
    # they sit at the end of a line, not at the start of one.
    markdown = re.sub(r"^<!--[\s\S]*?-->\n", "", markdown, flags=re.M)
    parts = sections(markdown)

    header = parts["_header"]
    name = re.search(r"^# (.+)$", header, flags=re.M)[1]
    contact = dict(re.findall(r"^- \*\*(.+?):\*\* (.+)$", header, flags=re.M))

    skills = [
        (label, split_items(items))
        for label, items in re.findall(r"^- \*\*(.+?):\*\* (.+)$", parts["Key Skills"], flags=re.M)
    ]

    experience: list[Company] = []
    earlier: list[EarlierRole] = []
    for heading, body in subsections(parts["Experience"], "###"):
        if heading == "Earlier Experience":
            for match in re.finditer(
                r"^- \*\*(.+?)\*\* — (.+) \(([^()]+)\): (.+)$", body, flags=re.M
            ):
                earlier.append(EarlierRole(*match.groups()))
            continue
        company_name, _, single_title = heading.partition(" — ")
        company_range = re.search(r"^\*\*(.+)\*\*$", body, flags=re.M)[1]
        note = re.search(r"^\*([^*].+)\*$", body, flags=re.M)
        company = Company(company_name, company_range, note[1] if note else None)
        roles = subsections(body, "####")
        if roles:
            for role_heading, role_body in roles:
                title, role_range = re.match(r"^(.+) \(([^()]+)\)$", role_heading).groups()
                company.roles.append(Role(title, role_range, parse_bullets(role_body)))
        else:
            company.roles.append(Role(single_title, company_range, parse_bullets(body)))
        experience.append(company)

    education = re.search(r"^\*\*(.+?)\*\* — (.+)$", parts["Education"], flags=re.M)
    interests = parts["Interests"].strip().split(" · ")

    variants = []
    for heading, body in subsections(parts["Resume Variants"], "###"):
        label, variant_id = re.match(r"^(.+) \((\w+)\)$", heading).groups()
        headline = re.search(r"^\*\*Headline:\*\* (.+)$", body, flags=re.M)[1]
        paragraphs = [p.strip() for p in body.split("\n\n") if p.strip()]
        summary = next(p for p in paragraphs if not p.startswith("**"))
        variants.append(Variant(variant_id, label, headline, summary))

    return Resume(
        name, contact, skills, experience, earlier, education.group(1, 2), interests, variants
    )


# ---------------------------------------------------------------------------
# python-docx helpers (things the high-level API doesn't expose)
# ---------------------------------------------------------------------------


def set_cell_shading(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_margins(cell, top=0, bottom=0, left=0, right=0) -> None:
    margins = OxmlElement("w:tcMar")
    for side, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(int(value.twips if hasattr(value, "twips") else value)))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    cell._tc.get_or_add_tcPr().append(margins)


def set_cell_border(cell, side: str, color: str, size_eighths: int = 6) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    edge = OxmlElement(f"w:{side}")
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), str(size_eighths))
    edge.set(qn("w:color"), color)
    borders.append(edge)


def fix_table_layout(table, widths) -> None:
    """Fixed column widths: Word otherwise autofits and ignores them."""
    widths = [Emu(int(width)) for width in widths]
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=True):
            cell.width = width
    grid = table._tbl.tblGrid
    for col, width in zip(grid.findall(qn("w:gridCol")), widths, strict=True):
        col.set(qn("w:w"), str(int(width.twips)))


def keep_row_together(row) -> None:
    cant_split = OxmlElement("w:cantSplit")
    row._tr.get_or_add_trPr().append(cant_split)


def band(container, width):
    """A one-cell shaded table: one continuous band with real inner padding.

    Paragraph shading can't do this. Word only merges shaded paragraphs into
    one box when their indents match, so indented sub-bullets broke out into
    their own offset box, and its border "padding" didn't inset the text.
    """
    table = container.add_table(rows=1, cols=1)
    fix_table_layout(table, [width])
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, BAND)
    set_cell_margins(cell, top=Pt(5), bottom=Pt(5), left=Pt(9), right=Pt(9))
    cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
    # _Cell.add_table appends the empty paragraph Word requires after a
    # table; pin its height so it becomes the gap below the band.
    gap = container.paragraphs[-1].paragraph_format
    gap.space_before = gap.space_after = Pt(0)
    gap.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    gap.line_spacing = Pt(3)
    return cell


def bottom_rule(paragraph, color: str = RULE) -> None:
    borders = OxmlElement("w:pBdr")
    edge = OxmlElement("w:bottom")
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), "6")
    edge.set(qn("w:space"), "1")
    edge.set(qn("w:color"), color)
    borders.append(edge)
    paragraph._p.get_or_add_pPr().append(borders)


def letter_spacing(run, twentieths_pt: int) -> None:
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), str(twentieths_pt))
    run._r.get_or_add_rPr().append(spacing)


def add_run(
    paragraph,
    text: str,
    *,
    font: str = BODY_FONT,
    size: float = 9.5,
    color: str = INK,
    bold: bool = False,
    italic: bool = False,
    caps: bool = False,
):
    run = paragraph.add_run(text)
    run.font.name = font
    run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.bold = bold
    run.font.italic = italic
    run.font.all_caps = caps
    return run


def add_hyperlink(paragraph, url: str, text: str, **run_kwargs) -> None:
    rel_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = add_run(paragraph, text, **run_kwargs)
    link.append(run._r)
    paragraph._p.append(link)


def paragraph(container, *, before: float = 0, after: float = 0):
    # No keep-with-next: inside a table cell Word applies it to the whole row,
    # which pushes the single-row body table off the first page.
    p = container.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.line_spacing = 1.0
    return p


# ---------------------------------------------------------------------------
# Document layout (see docs/style-guide.md, "Resume PDF template")
# ---------------------------------------------------------------------------


def section_heading(container, text: str, *, first: bool = False) -> None:
    p = paragraph(container, before=0 if first else 10, after=3)
    run = add_run(p, text, font=DISPLAY_FONT, size=11, color=BLUE, caps=True)
    letter_spacing(run, 10)


def bullet_paragraph(container, text: str, *, level: int = 0) -> None:
    indent = Inches(0.14 + 0.18 * level)
    p = paragraph(container, after=0.6)
    p.paragraph_format.left_indent = indent
    p.paragraph_format.first_line_indent = -Inches(0.13)
    p.paragraph_format.tab_stops.add_tab_stop(indent)
    p.paragraph_format.keep_together = True
    add_run(p, ("•" if level == 0 else "–") + "\t", color=BLUE if level == 0 else MUTED, size=9)
    add_run(p, text, size=9 if level == 0 else 8.6)


def build_header(doc, resume: Resume, variant: Variant) -> None:
    widths = [Inches(3.75), Inches(0.42), CONTENT_WIDTH - Inches(3.75) - Inches(0.42)]
    table = doc.add_table(rows=1, cols=3)
    fix_table_layout(table, widths)
    row = table.rows[0]
    row.height = Inches(1.0)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    keep_row_together(row)
    left, middle, right = row.cells

    for cell in (left, right):
        set_cell_shading(cell, BLOCK)
        set_cell_margins(cell, top=Pt(10), bottom=Pt(10), left=Pt(16), right=Pt(12))
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(middle, left=Pt(9), right=Pt(9))
    middle.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    first, _, last = resume.name.partition(" ")
    name_p = left.paragraphs[0]
    add_run(name_p, first.upper() + " ", font=DISPLAY_FONT, size=21, bold=True)
    add_run(name_p, last.upper(), font=DISPLAY_FONT, size=21, color=BLUE)
    title_p = paragraph(left, before=2)
    letter_spacing(add_run(title_p, variant.headline, font=DISPLAY_FONT, size=9.5, caps=True), 12)

    # Signature detail from the reference: three small blocks between the two header panels.
    squares = middle.add_table(rows=5, cols=1)
    fix_table_layout(squares, [Inches(0.42) - Pt(18)])
    for index, square_row in enumerate(squares.rows):
        square_row.height = Pt(14 if index % 2 == 0 else 8)
        square_row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        if index % 2 == 0:
            set_cell_shading(square_row.cells[0], BLOCK_SOFT)
    middle.paragraphs[0].paragraph_format.space_after = Pt(0)
    middle.paragraphs[0].runs or add_run(middle.paragraphs[0], "", size=1)

    contact = resume.contact
    links = {
        "Email": f"mailto:{contact['Email']}",
        "LinkedIn": contact["LinkedIn"],
        "GitHub": contact["GitHub"],
        "Website": contact["Website"],
    }
    first_line = True
    for label in ("Location", "Email", "LinkedIn", "GitHub", "Website"):
        value = contact[label]
        p = right.paragraphs[0] if first_line else paragraph(right, before=1.5)
        first_line = False
        add_run(p, f"{label}: ", size=8.4, bold=True)
        display = re.sub(r"^https?://(www\.)?", "", value)
        if label in links:
            add_hyperlink(p, links[label], display, size=8.4, color=INK)
        else:
            add_run(p, display, size=8.4)

    rule = paragraph(doc, before=6, after=8)
    bottom_rule(rule)
    add_run(rule, "", size=2)


def build_experience(container, resume: Resume, variant: Variant) -> None:
    section_heading(container, "Experience")
    for index, company in enumerate(resume.experience):
        single_role = len(company.roles) == 1 and company.roles[0].range == company.range
        company_p = paragraph(container, before=0 if index == 0 else 7, after=1)
        add_run(company_p, company.name, size=9.8, bold=True)
        if not single_role:
            meta = paragraph(container, after=2)
            details = [company.range] + ([company.note] if company.note else [])
            add_run(meta, " · ".join(details), size=8.4, color=MUTED)
        elif company.note:
            meta = paragraph(container, after=2)
            add_run(meta, company.note, size=8.4, color=MUTED)

        for role in company.roles:
            bullets = [b for b in role.bullets if variant.id in b.audiences]
            title_p = paragraph(container, before=3)
            add_run(title_p, role.title, font=DISPLAY_FONT, size=10.2, color=BLUE)
            range_p = paragraph(container, after=3 if bullets else 2)
            add_run(range_p, role.range, size=8.4, color=MUTED)
            if not bullets:
                continue
            box = band(container, MAIN_CONTENT_WIDTH)
            for bullet in bullets:
                bullet_paragraph(box, bullet.text)
                for item in bullet.sub_items:
                    bullet_paragraph(box, item, level=1)



def build_sidebar(container, resume: Resume) -> None:
    section_heading(container, "Skills", first=True)
    for label, items in resume.skills:
        label_p = paragraph(container, before=3, after=0.5)
        add_run(label_p, label, size=8.6, bold=True)
        items_p = paragraph(container, after=2)
        add_run(items_p, " · ".join(items), size=8.4, color=MUTED)

    section_heading(container, "Education")
    school, degree = resume.education
    p = paragraph(container, after=0)
    add_run(p, school, size=8.8, bold=True)
    p = paragraph(container)
    add_run(p, degree, size=8.4, color=MUTED)

    # Earlier roles live in the sidebar so they fill its otherwise empty
    # continuation on page 2 instead of lengthening the main column.
    section_heading(container, "Earlier Experience")
    for role in resume.earlier:
        p = paragraph(container, before=3)
        add_run(p, role.company, size=8.6, bold=True)
        p = paragraph(container)
        add_run(p, role.title, size=8.4, color=BLUE)
        p = paragraph(container, after=1)
        add_run(p, role.range, size=8, color=MUTED)
        p = paragraph(container, after=2)
        add_run(p, role.description, size=8.4, color=MUTED)

    section_heading(container, "Interests")
    for interest in resume.interests:
        p = paragraph(container, after=1.5)
        add_run(p, interest, size=8.4, color=MUTED)


def build_document(resume: Resume, variant: Variant) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width, section.page_height = PAGE_WIDTH, PAGE_HEIGHT
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, side, MARGIN)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(0)

    props = doc.core_properties
    props.title = f"{resume.name} — {variant.headline}"
    props.author = resume.name
    props.subject = f"Resume ({variant.label})"
    props.keywords = "resume, platform engineering, DevOps"

    build_header(doc, resume, variant)

    body = doc.add_table(rows=1, cols=2)
    fix_table_layout(body, [MAIN_WIDTH, CONTENT_WIDTH - MAIN_WIDTH])
    main, side = body.rows[0].cells
    set_cell_margins(main, left=0, right=MAIN_GUTTER)
    set_cell_margins(side, left=Pt(14), right=0)
    set_cell_border(side, "left", RULE_STRONG)

    main.paragraphs[0]._p.getparent().remove(main.paragraphs[0]._p)
    section_heading(main, "Profile Summary", first=True)
    summary = paragraph(band(main, MAIN_CONTENT_WIDTH))
    add_run(summary, variant.summary, size=9)
    build_experience(main, resume, variant)

    side.paragraphs[0]._p.getparent().remove(side.paragraphs[0]._p)
    build_sidebar(side, resume)

    # Word requires a paragraph after a final table; keep it invisible.
    tail = paragraph(doc)
    add_run(tail, "", size=1)
    return doc


# ---------------------------------------------------------------------------
# PDF conversion and manifest
# ---------------------------------------------------------------------------


def convert_with_word(docx_path: Path, pdf_path: Path) -> None:
    script = f'''
    tell application "Microsoft Word"
        open POSIX file "{docx_path}"
        set theDoc to active document
        save as theDoc file name "{pdf_path}" file format format PDF
        close theDoc saving no
    end tell
    '''
    subprocess.run(["osascript", "-e", script], check=True, capture_output=True, text=True)


def main() -> int:
    if sys.platform != "darwin" or not Path("/Applications/Microsoft Word.app").exists():
        print("pnpm resumes needs macOS with Microsoft Word installed.", file=sys.stderr)
        return 1

    source_bytes = SOURCE.read_bytes()
    resume = parse(source_bytes.decode("utf-8"))
    if {v.id for v in resume.variants} != set(FILE_SLUGS):
        print(f"Expected variants {sorted(FILE_SLUGS)} in {SOURCE.name}", file=sys.stderr)
        return 1

    DOCX_DIR.mkdir(parents=True, exist_ok=True)
    PDF_DIR.mkdir(parents=True, exist_ok=True)
    WORD_WORKDIR.mkdir(parents=True, exist_ok=True)

    editions = []
    for variant in resume.variants:
        stem = f"carlos-barajas-resume-{FILE_SLUGS[variant.id]}"
        docx_path = DOCX_DIR / f"{stem}.docx"
        build_document(resume, variant).save(docx_path)

        word_docx = WORD_WORKDIR / docx_path.name
        word_pdf = WORD_WORKDIR / f"{stem}.pdf"
        shutil.copyfile(docx_path, word_docx)
        word_pdf.unlink(missing_ok=True)
        convert_with_word(word_docx, word_pdf)
        shutil.copyfile(word_pdf, PDF_DIR / word_pdf.name)

        pages = len(PdfReader(PDF_DIR / word_pdf.name).pages)
        editions.append(
            {
                "id": variant.id,
                "label": variant.label,
                "headline": variant.headline,
                "file": f"/resumes/{word_pdf.name}",
                "pages": pages,
            }
        )
        print(f"{variant.label:24s} {pages} page(s)  public/resumes/{word_pdf.name}")

    manifest = {
        "source": SOURCE.name,
        "sourceSha256": hashlib.sha256(source_bytes).hexdigest(),
        "generated": date.today().isoformat(),
        "editions": editions,
    }
    MANIFEST.write_text(json.dumps(manifest, indent=2) + "\n")
    print(f"Wrote {MANIFEST.relative_to(REPO)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
